# QwizMate | SENG 564 | Fall 2026
# Author: Nick Moore

import io
from typing import List, Tuple

import pymupdf as fitz
from docx import Document
from pptx import Presentation
from pypdf import PdfReader

# Below this average character count per page, a PDF is treated as scanned/
# handwritten rather than text-based, and its pages are rasterized to images
# for the vision model to read directly instead of relying on extracted text.
MIN_CHARS_PER_PAGE_FOR_TEXT = 40


def extract_pdf(data: bytes) -> Tuple[str, List[bytes]]:
    """Returns (extracted_text, page_images_png).

    Text-based PDFs (typed slides exported to PDF, etc.) return their text
    with no images. Scanned or handwritten PDFs return no text and instead
    return each page rendered as a PNG, so the vision model can read them
    the same way it reads an uploaded photo.
    """
    reader = PdfReader(io.BytesIO(data))
    texts = [page.extract_text() or "" for page in reader.pages]
    total_chars = sum(len(t) for t in texts)
    avg_chars = total_chars / max(len(texts), 1)

    if avg_chars >= MIN_CHARS_PER_PAGE_FOR_TEXT:
        return "\n\n".join(texts), []

    doc = fitz.open(stream=data, filetype="pdf")
    images = [page.get_pixmap(dpi=150).tobytes("png") for page in doc]
    return "", images


def extract_pptx(data: bytes) -> str:
    prs = Presentation(io.BytesIO(data))
    chunks = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = "\n".join(p.text for p in shape.text_frame.paragraphs)
                if text.strip():
                    chunks.append(text)
    return "\n\n".join(chunks)


def extract_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    chunks = [p.text for p in doc.paragraphs if p.text.strip()]
    # Study notes often keep definitions/comparisons in tables, which
    # doc.paragraphs doesn't include.
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                chunks.append(" | ".join(cells))
    return "\n\n".join(chunks)
